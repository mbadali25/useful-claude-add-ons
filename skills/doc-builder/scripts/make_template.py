"""
Regenerate a brand pack's sop_template.docx -- the blank house template
build_sop.py builds on -- from a conforming master.

The template is derived from the canonical master rather than built from
scratch, so it inherits the exact Normal style (Franklin Gothic Book), the
theme, the section margins (0.8/0.6/0.5/0.5) and the three-cell footer table
without those having to be re-declared in code.

What this script does:
  1. Open the reference master.
  2. Delete every body block element except the trailing <w:sectPr>.
  3. Drop the image relationships so the template does not carry ~700 KB of
     the reference SOP's screenshots.

Run this only if the house style changes or the template is lost. The
reference master must be named; nothing here guesses which document is the
canonical one:

    python make_template.py --reference "C:\\repos\\OnboardingSOPs\\sops_new\\How to access LinkedIn Learning.docx"
    python make_template.py --reference "<master.docx>" --output "<brand>\\assets\\sop_template.docx"

Without --output the template path in the active brand pack is overwritten;
a pack with no template (the neutral pack) needs --output.

Third-party requirement: python-docx.
"""
from __future__ import annotations

import argparse
import os

import docx
from docx.oxml.ns import qn

import resolve_brand


def make_template(reference: str, output: str) -> str:
    doc = docx.Document(reference)
    body = doc.element.body

    # 1. strip body content, keeping the section properties
    sectPr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sectPr:
            body.remove(child)

    # 2. drop image parts (rels not referenced by the now-empty body)
    for rId, rel in list(doc.part.rels.items()):
        if "image" in rel.reltype:
            doc.part.drop_rel(rId)

    doc.save(output)
    return output


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--reference", required=True,
                    help="conforming SOP master to derive the template from")
    ap.add_argument("--output", help="where to write (default: the brand pack's template path)")
    resolve_brand.add_brand_argument(ap)
    args = ap.parse_args()

    output = args.output
    if not output:
        brand = resolve_brand.resolve(args.brand)
        output = brand.template
        if not output:
            ap.error("the %s brand pack has no template path; pass --output" % brand.name)

    ref = os.path.normpath(args.reference)
    out = make_template(ref, os.path.normpath(output))
    size_kb = os.path.getsize(out) / 1024

    d = docx.Document(out)
    s = d.sections[0]
    print("Reference : %s" % ref)
    print("Template  : %s (%.0f KB)" % (out, size_kb))
    print("  Normal font : %s" % d.styles["Normal"].font.name)
    print("  Margins     : T%.2f B%.2f L%.2f R%.2f in"
          % (s.top_margin.inches, s.bottom_margin.inches,
             s.left_margin.inches, s.right_margin.inches))
    print("  Body blocks : %d (expect 0)" % len(d.paragraphs))
    print("  Footer table: %s (expect True)" % bool(s.footer.tables))
    print("  Image parts : %d (expect 0)"
          % sum(1 for r in d.part.rels.values() if "image" in r.reltype))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
