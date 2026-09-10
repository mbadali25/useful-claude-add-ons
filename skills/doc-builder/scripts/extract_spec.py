"""
Reverse-engineer a JSON build spec from an existing SOP .docx master built on
the active brand pack's house template.

Why: build_sop.py builds SOPs from JSON specs, but most production masters
predate it and have none. Those can only be hand-edited in Word today --
exactly the failure mode this toolchain exists to prevent. Council decision: specs are extracted
JUST-IN-TIME, one master at a time, when that master is next actually edited
-- not bulk-generated for all 18 and committed unverified. This tool has two
jobs to match that policy:

  1. Extract one master well (mode 1).
  2. Honestly report the whole set's extractability, so the just-in-time
     work has a queue (mode 2).

Usage
-----
    # Extract one master to <brand specs_dir>\\<slug>.json and
    # <brand assets_dir>\\<slug>\\ (defaults):
    python extract_spec.py "C:\\repos\\OnboardingSOPs\\sops_new\\SafeSend SOP.docx"

    # Override slug / output locations, e.g. to extract to a scratch location
    # for review before it's the real spec:
    python extract_spec.py "<master.docx>" --slug my-slug \\
        --out <temp>\\my-slug.json --assets-out <temp>\\my-slug-assets

    # Refuse to overwrite an existing spec or asset file unless told to:
    python extract_spec.py "<master.docx>" --force

    # Report on the whole set -- the queue for the just-in-time policy:
    python extract_spec.py --report

Classification rules are NOT reinvented here. Block-type classification
(heading / step / bullet / caption / image / spacer / tip / para) is imported
directly from check_conformance._classify_para -- the working classifier
that already gates every master's conformance -- so this tool can never
diverge from what the checker considers "correct". Constants (fonts, colours,
sizes, the accent/border colours, the effectExtent floor) come from
build_sop.py, configured for the active brand, never re-hardcoded here.

Honesty over optimism (the whole point of this tool): every paragraph and
every run is accounted for. Anything this extractor cannot express in the
spec format -- italic/underline/strike/highlight/caps runs, non-house
hyperlink styling, a second body table, a floating image, a Word-side crop,
a footnote, an unresolved link target, a marker shape it doesn't recognise --
is recorded in the extracted spec's own "_warnings" array, not silently
dropped. After building the spec in memory, this tool re-runs
check_conformance's own spec/master diff (_spec_blocks + _diff_blocks)
against the source master as a self-check: if the spec it just built would
not reproduce the same block sequence the checker sees in the master, that
is reported too, never swallowed.

Requires: python-docx (the same dependency as the rest of this toolchain).
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
from collections import Counter

import docx
from docx.oxml.ns import qn
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.shared import Inches

import check_conformance as C
import build_sop as S
import resolve_brand

# Bound by main(); module-level so the helpers below need no plumbing change.
BRAND = None

EMU_PER_INCH = int(Inches(1))


# --------------------------------------------------------------------------
# Small text helpers
# --------------------------------------------------------------------------
def _specs_dir():
    d = BRAND.specs_dir if BRAND else None
    if not d:
        raise SystemExit("The %s brand pack has no specs directory; pass --out."
                         % (BRAND.name if BRAND else "active"))
    return d


def _slugify(text, max_words=10):
    """Best-effort filename/dir slug: lowercase, non-alnum -> '-', trimmed."""
    text = re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")
    words = [w for w in text.split("-") if w]
    return "-".join(words[:max_words]) or "sop"


def _to_spec_relpath(target_abspath, spec_dir):
    """A path in the emitted spec is always relative to the spec file itself
    (per TEMPLATE-SPEC.md / the qrg-archiving.json ground truth), and always
    forward-slashed so the JSON is portable off Windows."""
    rel = os.path.relpath(target_abspath, start=spec_dir)
    return rel.replace(os.sep, "/")


_EXT_BY_CONTENT_TYPE = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/tiff": ".tif",
    "image/x-emf": ".emf",
    "image/x-wmf": ".wmf",
}


def _ext_for_part(part):
    ext = _EXT_BY_CONTENT_TYPE.get(part.content_type)
    if ext:
        return ext
    return os.path.splitext(str(part.partname))[1] or ".bin"


# --------------------------------------------------------------------------
# Warning ledger -- every entry is something lossy, guessed, or uncertain.
# Tagged "[category] context: message" so --report can tally by category
# without re-parsing free text. This IS the honesty mechanism the task asks
# for: nothing below is ever swallowed, only ever appended here.
# --------------------------------------------------------------------------
class Ledger(object):
    def __init__(self):
        self.warnings = []

    def add(self, category, ctx, message):
        self.warnings.append("[%s] %s: %s" % (category, ctx, message))

    def category_counts(self):
        counts = Counter()
        for w in self.warnings:
            m = re.match(r"^\[([\w-]+)\]", w)
            if m:
                counts[m.group(1)] += 1
        return counts


# --------------------------------------------------------------------------
# Inline markup: docx runs/hyperlinks (in document order) -> "**bold**" /
# "[label](url)" text, the exact syntax _split_spans() in build_sop.py
# parses. This is that function's inverse.
# --------------------------------------------------------------------------
def _rgb(run):
    return C._rgb(run)


def _flag_exotic_run_formatting(run, ctx, ledger, check_underline=True,
                                italic_is_house=False):
    """Any run formatting the spec's "**bold**"-only markup cannot express.
    Text is still kept (best-effort) -- only the styling is lost -- but every
    occurrence is logged, never silently dropped.

    italic_is_house: pass True inside a caption. Captions are italic BY HOUSE
    STYLE -- measured 2026-09-08 as 47 italic caption paragraphs across 17
    masters and zero plain -- and build_sop.caption() now applies it, so an
    italic caption run is reproduced faithfully on rebuild and is not a loss.
    Reporting it would have buried the 3 real inline-italic runs in the set
    under 103 false ones.
    """
    f = run.font
    bad = []
    if f.italic and not italic_is_house:
        bad.append("italic")
    if check_underline and f.underline:
        bad.append("underline")
    if f.strike or f.double_strike:
        bad.append("strike")
    if f.highlight_color:
        bad.append("highlight=%s" % f.highlight_color)
    if f.all_caps:
        bad.append("all_caps")
    if f.small_caps:
        bad.append("small_caps")
    if f.subscript:
        bad.append("subscript")
    if f.superscript:
        bad.append("superscript")
    if bad:
        text = (run.text or "")[:40]
        ledger.add("formatting", ctx,
                   "run %r has %s -- not expressible as **bold**/[link] markup, "
                   "DROPPED (text kept, styling lost on rebuild)"
                   % (text, ", ".join(bad)))


def _run_markup(run, ctx, ledger, expect_bold, italic_is_house=False):
    """One plain (non-hyperlink) run -> its markup chunk."""
    text = run.text or ""
    if text == "":
        return ""
    _flag_exotic_run_formatting(run, ctx, ledger,
                                italic_is_house=italic_is_house)
    if "**" in text:
        ledger.add("literal-markup", ctx,
                    "run %r contains a literal '**' -- will be misread as a bold "
                    "delimiter when this spec is rebuilt" % text[:40])
    if S._LINK_RE.search(text):
        ledger.add("literal-markup", ctx,
                    "run %r contains a '[..](..)'  -shaped substring that is not "
                    "an actual hyperlink here -- will be misread as a link when "
                    "this spec is rebuilt" % text[:40])
    bold = bool(run.bold)
    if bold and not expect_bold:
        return "**%s**" % text
    return text


def _hyperlink_label_markup(hyperlink, ctx, ledger):
    """A hyperlink's label may itself contain **bold** spans (build_sop.py's
    _split_spans supports this -- "a link's label may itself contain **bold**
    (but not another link)"). Build that label markup from its runs."""
    parts = []
    for r in hyperlink.runs:
        text = r.text or ""
        if not text:
            continue
        _flag_exotic_run_formatting(r, ctx, ledger, check_underline=False)
        if "[" in text or "]" in text:
            ledger.add("literal-markup", ctx,
                        "hyperlink label run %r contains '[' or ']' -- will break "
                        "the [label](url) delimiter on rebuild" % text[:40])
        parts.append("**%s**" % text if r.bold else text)
    return "".join(parts)


def _check_hyperlink_style(hyperlink, ctx, ledger):
    """House hyperlink style is LINK_BLUE + single underline (measured; see
    build_sop.LINK_BLUE and _add_hyperlink). A rebuild always renders house
    style regardless of the source, so any deviation here is a real, lossy
    conversion worth flagging -- not just cosmetic noise."""
    for r in hyperlink.runs:
        colour = _rgb(r)
        if colour is not None and colour.upper() != S.LINK_BLUE:
            ledger.add("hyperlink-style", ctx,
                       "hyperlink %r run colour is %s, house is %s -- rebuild "
                       "will use house colour regardless (original look lost)"
                       % (hyperlink.text[:40], colour, S.LINK_BLUE))
        if not r.font.underline:
            ledger.add("hyperlink-style", ctx,
                       "hyperlink %r is not underlined in the source -- rebuild "
                       "always underlines (original look lost)" % hyperlink.text[:40])


def _items_to_markup(items, ctx, ledger, expect_bold=False,
                     italic_is_house=False):
    """items: a list from Paragraph.iter_inner_content() (Run and Hyperlink
    objects, in true document order -- this is what makes hyperlinks visible
    at all; Paragraph.runs silently skips over anything inside a
    <w:hyperlink>). expect_bold=True means the whole block is force-bold by
    the generator (heading/tip) so per-run bold carries no information and is
    not re-encoded as "**" -- that would be inert noise, not a defect."""
    parts = []
    for item in items:
        if isinstance(item, Hyperlink):
            url = None
            try:
                url = item.address
            except Exception:
                url = None
            label_md = _hyperlink_label_markup(item, ctx, ledger)
            if not url:
                ledger.add("unresolved-link", ctx,
                           "hyperlink %r has no resolvable target (r:id missing or "
                           "dead relationship) -- emitted as plain text, the LINK "
                           "ITSELF IS LOST" % item.text[:40])
                parts.append(label_md.replace("**", ""))
                continue
            _check_hyperlink_style(item, ctx, ledger)
            parts.append("[%s](%s)" % (label_md, url))
        elif isinstance(item, Run):
            parts.append(_run_markup(item, ctx, ledger, expect_bold,
                                     italic_is_house=italic_is_house))
        else:
            ledger.add("structure", ctx,
                       "paragraph contains a %r inline-content item this extractor "
                       "does not handle -- its text, if any, is DROPPED"
                       % type(item).__name__)
    return "".join(parts)


def _strip_marker_items(items, marker_ok, ctx, ledger):
    """Drop the leading marker (+ tab) that build_sop.py's step()/bullet()
    emit ahead of the block's own text, so they are not doubled into the
    extracted text. marker_ok(stripped_text) tests the first item's text
    (marker-only, tab stripped) against the expected marker shape.

    Two shapes are seen across the real masters: the generator's own
    marker-run-then-separate-tab-run ("1." then "\\t"), and a hand-authored
    master with the marker and its tab combined into one run ("1.\\t")."""
    if items and isinstance(items[0], Run):
        raw = items[0].text or ""
        if marker_ok(raw.strip()):
            if raw.endswith("\t"):
                return items[1:]  # marker + tab combined in one run
            if (len(items) >= 2 and isinstance(items[1], Run)
                    and (items[1].text or "") == "\t"):
                return items[2:]  # marker run, then a separate tab run
    ledger.add("marker-shape", ctx,
               "expected marker (+ tab) prefix not found in either shape this "
               "extractor recognises -- the marker text may be duplicated in "
               "the extracted text")
    return items


# --------------------------------------------------------------------------
# Title banner
# --------------------------------------------------------------------------
def _extract_banner(doc, ledger):
    """Title/subtitle live in the 1-row/2-col banner table's last cell (see
    SopBuilder._add_title_banner). Returns (title, subtitle)."""
    if not doc.tables:
        ledger.add("banner", "banner", "no title banner table -- title/subtitle "
                   "extraction FAILED, both left blank")
        return "", ""
    tbl = doc.tables[0]
    fills = re.findall(r'<w:shd[^>]*w:fill="([0-9A-Fa-f]{6})"', tbl._tbl.xml)
    if S.ACCENT_RED not in [f.upper() for f in fills]:
        ledger.add("banner", "banner",
                   "accent bar cell is not shaded %s (found %s) -- rebuild always "
                   "uses house accent colour regardless" % (S.ACCENT_RED, fills))
    if not tbl.rows or len(tbl.rows[0].cells) < 2:
        ledger.add("banner", "banner", "banner table does not have the expected "
                   "2 cells -- title/subtitle extraction FAILED")
        return "", ""
    text_cell = tbl.rows[0].cells[-1]
    texts = [p.text for p in text_cell.paragraphs if p.text.strip()]
    title = texts[0] if texts else ""
    subtitle = texts[1] if len(texts) > 1 else ""
    if not title:
        ledger.add("banner", "banner", "banner cell has no title text")
    if len(texts) > 2:
        ledger.add("banner", "banner",
                   "banner cell has %d non-blank paragraphs, expected 1 (title) or "
                   "2 (title+subtitle) -- extra text %r is DROPPED"
                   % (len(texts), texts[2:]))
    return title, subtitle


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------
def _extract_image(para, doc, ctx, ledger):
    """One image-bearing paragraph -> (width_in, blob, ext, name_hint) or None
    if it turns out to carry no extractable picture at all. Handles both the
    normal inline case and (defensively -- the survey found none in this set)
    a floating/anchored one."""
    p_el = para._p
    inlines = p_el.findall(".//" + qn("wp:inline"))
    anchors = p_el.findall(".//" + qn("wp:anchor"))
    vml_picts = p_el.findall(".//" + qn("w:pict"))

    if vml_picts:
        ledger.add("legacy-image", ctx,
                   "%d legacy VML <w:pict> image(s) found -- not extracted (this "
                   "format predates the DrawingML pictures every master in the "
                   "current survey uses)" % len(vml_picts))

    if anchors:
        ledger.add("anchor-image", ctx,
                   "%d floating/anchored image(s) (wp:anchor) found -- this "
                   "contradicts the survey finding that no master uses floating "
                   "images. Extracted as a plain inline centered picture like "
                   "every other image block; the original floating position and "
                   "text-wrap behaviour CANNOT be expressed by the image block "
                   "type and is LOST" % len(anchors))
        inlines = inlines + anchors  # best-effort: still try to pull pixels/width

    if not inlines:
        return None
    if len(inlines) > 1:
        ledger.add("multi-image", ctx,
                   "%d images in one paragraph -- only the first is extracted, "
                   "the rest are DROPPED" % len(inlines))

    inline = inlines[0]
    extent = inline.find(qn("wp:extent"))
    if extent is None:
        ledger.add("structure", ctx, "image has no wp:extent -- width unknown, "
                   "skipped entirely")
        return None
    cx = int(extent.get("cx"))
    width_in = round(cx / EMU_PER_INCH, 2)

    blip = inline.find(".//" + qn("a:blip"))
    rid = blip.get(qn("r:embed")) if blip is not None else None
    if rid is None:
        ledger.add("structure", ctx, "image has no r:embed relationship -- cannot "
                   "extract its bytes, skipped entirely")
        return None
    rel = doc.part.rels.get(rid)
    if rel is None:
        ledger.add("structure", ctx, "image r:embed %r has no matching "
                   "relationship -- cannot extract its bytes, skipped entirely" % rid)
        return None
    target_part = rel.target_part
    blob = target_part.blob
    ext = _ext_for_part(target_part)

    if inline.find(".//" + qn("a:srcRect")) is not None:
        ledger.add("crop", ctx,
                   "image has an a:srcRect (Word-side cropping) -- this "
                   "contradicts the survey finding that no master crops images "
                   "this way. The extracted PNG is the FULL, UNCROPPED source "
                   "image; the crop is LOST")

    docPr = inline.find(".//" + qn("wp:docPr"))
    name_hint = None
    if docPr is not None:
        name_hint = docPr.get("descr") or docPr.get("name")

    return width_in, blob, ext, name_hint


# --------------------------------------------------------------------------
# One paragraph -> one spec block (dispatch by the reused classifier)
# --------------------------------------------------------------------------
_STEP_MARKER_RE = re.compile(r"^\d+\.$")


def _walk_document(doc, ctx_name, ledger):
    """Walk the body in true document order (mixing w:p / w:tbl / w:sectPr,
    not python-docx's separate .paragraphs/.tables lists, so a second body
    table -- which should not exist per the survey, but if it did would
    otherwise vanish silently) -> (title, subtitle, body_blocks, stats).

    stats is a collections.Counter of block-type counts for --report.
    """
    title, subtitle = _extract_banner(doc, ledger)

    stats = Counter()
    body_blocks = []
    children = list(doc.element.body)
    expected_step_no = 1
    tables_seen = 0
    i = 0
    while i < len(children):
        el = children[i]
        if el.tag == qn("w:tbl"):
            tables_seen += 1
            if not (tables_seen == 1 and i == 0):
                preview = re.sub(r"\s+", " ", "".join(
                    t.text or "" for t in el.findall(".//" + qn("w:t"))
                ))[:80]
                ledger.add("extra-table", "body position %d" % i,
                           "a %s table beyond the title banner -- this contradicts "
                           "the survey finding of at most 1 table per master. Its "
                           "content (preview: %r) is NOT representable by this spec "
                           "format and was NOT extracted" % (
                               "%dx%d" % (len(el.findall(qn("w:tr"))),
                                          len(el.find(qn("w:tr")).findall(qn("w:tc")))
                                          if el.find(qn("w:tr")) is not None else 0),
                               preview))
                stats["unrepresentable_table"] += 1
            i += 1
            continue
        if el.tag == qn("w:sectPr"):
            i += 1
            continue
        if el.tag != qn("w:p"):
            ledger.add("structure", "body position %d" % i,
                       "unrecognized top-level body element <%s> -- skipped, not "
                       "representable" % el.tag.split("}")[-1])
            stats["unrepresentable_element"] += 1
            i += 1
            continue

        para = Paragraph(el, doc)
        ctx = "%s para %d %r" % (ctx_name, i, para.text[:40])
        kind, _ = C._classify_para(para)

        if kind == "image":
            img = _extract_image(para, doc, ctx, ledger)
            stats["image"] += 1
            block = {"type": "image"}
            caption_text = ""
            if img is not None:
                width_in, blob, ext, name_hint = img
                block["_blob"] = blob
                block["_ext"] = ext
                block["_name_hint"] = name_hint
                block["width_in"] = width_in
            else:
                block["_blob"] = None
            # A caption is its own following paragraph in the rendered docx
            # (SopBuilder.image() calls self.caption() right after adding the
            # picture) but a FIELD of the image block in the spec.
            if i + 1 < len(children) and children[i + 1].tag == qn("w:p"):
                nxt = Paragraph(children[i + 1], doc)
                nkind, _ = C._classify_para(nxt)
                if nkind == "caption":
                    cap_items = list(nxt.iter_inner_content())
                    caption_text = _items_to_markup(cap_items, ctx + " caption", ledger,
                                                    italic_is_house=True)
                    i += 1  # consume it
            if caption_text:
                block["caption"] = caption_text
            body_blocks.append(block)
            i += 1
            continue

        if kind == "spacer":
            stats["spacer"] += 1
            body_blocks.append({"type": "spacer"})
            i += 1
            continue

        items = list(para.iter_inner_content())

        if kind == "heading":
            stats["heading"] += 1
            text = _items_to_markup(items, ctx, ledger, expect_bold=True)
            body_blocks.append({"type": "heading", "text": text})
            expected_step_no = 1
            i += 1
            continue

        if kind == "step":
            stats["step"] += 1
            content = _strip_marker_items(items, lambda t: bool(_STEP_MARKER_RE.match(t)),
                                          ctx, ledger)
            text = _items_to_markup(content, ctx, ledger, expect_bold=False)
            actual_no = None
            m = _STEP_MARKER_RE.match((items[0].text or "").strip()) if items else None
            if m:
                actual_no = int(items[0].text.strip().rstrip("."))
            block = {"type": "step", "text": text}
            if actual_no is not None and actual_no != expected_step_no:
                block["number"] = actual_no
                ledger.add("custom-number", ctx,
                           "step is numbered %d, auto-numbering would give %d -- "
                           "explicit \"number\" recorded" % (actual_no, expected_step_no))
                expected_step_no = actual_no + 1
            elif actual_no is not None:
                expected_step_no = actual_no + 1
            else:
                expected_step_no += 1
            body_blocks.append(block)
            i += 1
            continue

        if kind == "bullet":
            stats["bullet"] += 1
            content = _strip_marker_items(items, lambda t: t == S.BULLET_CHAR, ctx, ledger)
            text = _items_to_markup(content, ctx, ledger, expect_bold=False)
            body_blocks.append({"type": "bullet", "text": text})
            i += 1
            continue

        if kind == "caption":
            # An orphan caption-styled paragraph not immediately following an
            # image. Unusual (captions normally get folded into the image
            # block above) but build_sop.py's dispatch does support a
            # standalone "caption" block, so use that rather than inventing
            # anything or silently dropping the text.
            stats["caption"] += 1
            text = _items_to_markup(items, ctx, ledger,
                                    italic_is_house=True)
            ledger.add("orphan-caption", ctx,
                       "caption-styled paragraph does not immediately follow an "
                       "image -- emitted as a standalone \"caption\" block")
            body_blocks.append({"type": "caption", "text": text})
            i += 1
            continue

        if kind == "tip":
            stats["tip"] += 1
            text = _items_to_markup(items, ctx, ledger, expect_bold=True)
            body_blocks.append({"type": "tip", "text": text})
            i += 1
            continue

        # kind == "para"
        stats["para"] += 1
        text = _items_to_markup(items, ctx, ledger, expect_bold=False)
        body_blocks.append({"type": "para", "text": text})

        # Safety net for the one real blind spot in the reused classifier: a
        # heading-sized paragraph that fails the *font* half of the heading
        # test (e.g. a typo'd font name) silently falls through to here as an
        # ordinary paragraph. That is a property of the classifier this tool
        # is told to reuse verbatim, not something to special-case away -- but
        # it is exactly the kind of thing this task's honesty rule exists
        # for, so flag it without changing the classification.
        first_run = para.runs[0] if para.runs else None
        if first_run is not None:
            size_pt = first_run.font.size.pt if first_run.font.size else None
            if size_pt == S.HEADING_PT and first_run.font.name != S.HEADING_FONT:
                ledger.add("possible-heading", ctx,
                           "paragraph is %.1fpt (heading size) but font is %r, not "
                           "%r -- the classifier does not count this as a heading; "
                           "verify by eye" % (size_pt, first_run.font.name, S.HEADING_FONT))
        i += 1

    return title, subtitle, body_blocks, stats


# --------------------------------------------------------------------------
# Self-check: does the spec this tool just built reproduce the same block
# sequence check_conformance.py sees in the source master? This reuses
# the checker's OWN spec/master drift detector (_spec_blocks + _diff_blocks),
# so it can never diverge from what "correct" means to the tool that gates
# every master already. A diff here is either a real defect in this
# extractor, or a representational choice that needs calling out by name --
# never something to paper over.
# --------------------------------------------------------------------------
def _self_check(doc, spec_for_check, ledger):
    try:
        expected = C._spec_blocks(spec_for_check)
        actual = [C._classify_para(p) for p in doc.paragraphs]
        diff = C._diff_blocks("(freshly extracted spec)", expected, actual)
    except Exception as exc:  # never let the self-check crash extraction
        ledger.add("selfcheck", "self-check", "could not run: %s" % exc)
        return
    if diff:
        ledger.add("selfcheck", "self-check", diff)


# --------------------------------------------------------------------------
# Mode 1: extract one master
# --------------------------------------------------------------------------
def build_spec(master_path, slug, ledger):
    """Pure, side-effect-free extraction: opens the master, walks it, returns
    a spec dict (image blocks still carry "_blob"/"_ext"/"_name_hint" -- the
    caller decides whether/where to write those bytes) plus doc/stats for the
    self-check and --report. No file is written here."""
    doc = docx.Document(master_path)
    title, subtitle, body_blocks, stats = _walk_document(doc, slug, ledger)
    spec = {
        "title": title,
        "subtitle": subtitle,
        "output": None,  # filled in by the caller, which knows the spec path
        "body": body_blocks,
    }
    return doc, spec, stats


def _finalize_and_write(doc, spec, master_path, slug, out_path, assets_dir,
                        force, ledger):
    """Resolve image blocks' "_blob" payloads into real files + spec "path"
    fields, resolve "output", run the self-check, write everything, and
    return the list of files actually written."""
    spec_dir = os.path.dirname(out_path)
    spec["output"] = _to_spec_relpath(os.path.abspath(master_path), spec_dir)

    # Safety: a real spec (one that lives in specs/) must point "output" at the
    # master it reproduces -- that is the whole point of the field. But a spec
    # extracted to a scratch directory for testing must NOT, because building
    # it would resolve that path and silently overwrite the production master.
    # That happened once during this tool's own development. Redirect a
    # scratch extraction's output to a scratch .docx instead, and say so.
    _norm = lambda p: os.path.normcase(os.path.normpath(os.path.abspath(p)))
    # The brand's specs directory is only a DEFAULT. A pack with none (neutral)
    # must not stop an extraction whose --out was given explicitly; such an
    # extraction is simply always "scratch" and gets the redirect below.
    real_specs = BRAND.specs_dir if BRAND else None
    if real_specs is None or _norm(spec_dir) != _norm(real_specs):
        scratch_docx = os.path.splitext(os.path.basename(out_path))[0] + ".docx"
        spec["output"] = scratch_docx
        spec.setdefault("_warnings", []).append(
            "Extracted outside %s, so \"output\" was redirected to the local "
            "%r rather than the real master, to stop a rebuild overwriting "
            "production. Point it at the master deliberately if that is what "
            "you want." % (real_specs or "the brand's specs directory (this pack has none)",
                           scratch_docx))
        sys.stderr.write(
            "note: extracted outside specs/, so \"output\" points at %r rather\n"
            "      than the real master -- building this spec will not\n"
            "      overwrite production.\n" % scratch_docx)

    planned_images = []  # (abs_path, blob)
    seq = 0
    for blk in spec["body"]:
        if blk.get("type") != "image":
            continue
        seq += 1
        blob = blk.pop("_blob", None)
        ext = blk.pop("_ext", ".png")
        name_hint = blk.pop("_name_hint", None)
        if blob is None:
            # Extraction failed for this image (already logged); nothing to
            # write and no usable "path" -- record a placeholder rather than
            # emitting a spec the builder would crash on with a KeyError.
            blk["path"] = None
            ledger.add("image-missing", "image %d" % seq,
                       "could not be extracted -- \"path\" is null; this spec "
                       "CANNOT be built as-is")
            continue
        source_text = blk.get("caption") or name_hint or ""
        name_slug = _slugify(source_text, max_words=4) or "img"
        fname = "%02d-%s%s" % (seq, name_slug, ext)
        abs_path = os.path.join(assets_dir, fname)
        planned_images.append((abs_path, blob))
        blk["path"] = _to_spec_relpath(abs_path, spec_dir)

    # Self-check BEFORE writing anything -- a spec dict with the temporary
    # "_blob" keys stripped is what will actually be built, so check that one.
    _self_check(doc, spec, ledger)

    spec["_warnings"] = list(ledger.warnings)

    # Overwrite protection: check every planned write up front so a partial
    # write never happens because of a conflict discovered halfway through.
    conflicts = []
    if os.path.exists(out_path) and not force:
        conflicts.append(out_path)
    for abs_path, _blob in planned_images:
        if os.path.exists(abs_path) and not force:
            conflicts.append(abs_path)
    if conflicts:
        raise FileExistsError(
            "refusing to overwrite %d existing file(s) without --force:\n  %s"
            % (len(conflicts), "\n  ".join(conflicts)))

    os.makedirs(spec_dir, exist_ok=True)
    if planned_images:
        os.makedirs(assets_dir, exist_ok=True)
    for abs_path, blob in planned_images:
        with open(abs_path, "wb") as fh:
            fh.write(blob)

    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(spec, fh, indent=2, ensure_ascii=False)
        fh.write("\n")

    return [out_path] + [p for p, _ in planned_images]


def extract_one(master_path, slug=None, out_path=None, assets_dir=None, force=False):
    master_path = os.path.abspath(master_path)
    if not os.path.isfile(master_path):
        raise SystemExit("master not found: %s" % master_path)
    slug = slug or _slugify(os.path.splitext(os.path.basename(master_path))[0])
    out_path = os.path.abspath(out_path or os.path.join(_specs_dir(), slug + ".json"))
    assets_dir = os.path.abspath(assets_dir or os.path.join(BRAND.assets_dir, slug))

    ledger = Ledger()
    doc, spec, stats = build_spec(master_path, slug, ledger)
    written = _finalize_and_write(doc, spec, master_path, slug, out_path, assets_dir,
                                  force, ledger)

    print("Wrote: %s" % out_path)
    for p in written[1:]:
        print("Wrote: %s" % p)
    print("\nBlocks: %s" % ", ".join("%d %s" % (n, k) for k, n in sorted(stats.items())))
    print("Warnings: %d" % len(ledger.warnings))
    for w in ledger.warnings:
        print("  %s" % w)
    return out_path


# --------------------------------------------------------------------------
# Mode 2: report on the whole set
# --------------------------------------------------------------------------
_BLOCK_KINDS = ("heading", "step", "bullet", "tip", "para", "caption", "spacer", "image")


def report_all():
    paths = sorted(glob.glob(os.path.join(BRAND.require_masters_dir(), "*.docx")))
    paths = [p for p in paths if not os.path.basename(p).startswith("~$")]
    if not paths:
        print("No .docx masters found.", file=sys.stderr)
        return 2

    rows = []
    for path in paths:
        name = os.path.basename(path)
        ledger = Ledger()
        try:
            doc, spec, stats = build_spec(path, _slugify(os.path.splitext(name)[0]), ledger)
            # Strip the temporary blob payloads for the self-check spec, but
            # don't actually resolve/write image paths in report mode -- give
            # every image block a placeholder path so _spec_blocks() (which
            # only cares about type/text/caption) can still run.
            check_spec = json.loads(json.dumps({
                "title": spec["title"], "subtitle": spec["subtitle"],
                "output": "x", "body": [
                    {k: v for k, v in blk.items() if not k.startswith("_")}
                    for blk in spec["body"]
                ],
            }))
            for blk in check_spec["body"]:
                if blk.get("type") == "image":
                    blk.setdefault("path", "x")
            _self_check(doc, check_spec, ledger)
            n_links = sum(1 for blk in spec["body"]
                          for _ in S._LINK_RE.finditer(blk.get("text") or blk.get("caption") or ""))
            image_fail = sum(1 for blk in spec["body"]
                             if blk.get("type") == "image" and blk.get("_blob") is None)
        except Exception as exc:  # a corrupt/unreadable master must not stop the batch
            rows.append((name, None, None, None, str(exc)))
            continue
        rows.append((name, stats, n_links, image_fail, ledger))

    print("%-55s %-42s %5s %5s %6s" % ("MASTER", "BLOCKS", "IMGS", "LINKS", "ISSUES"))
    clean, needs_attention = [], []
    for row in rows:
        name = row[0]
        if row[1] is None:
            print("%-55s ERROR: %s" % (name[:55], row[4]))
            needs_attention.append((name, "could not be read: %s" % row[4]))
            continue
        _, stats, n_links, image_fail, ledger = row
        blocks_str = ", ".join("%d %s" % (stats[k], k) for k in _BLOCK_KINDS if stats.get(k))
        issues = len(ledger.warnings)
        print("%-55s %-42s %5d %5d %6d" % (
            name[:55], blocks_str[:42], stats.get("image", 0), n_links, issues))
        if issues == 0 and image_fail == 0:
            clean.append(name)
        else:
            cats = ledger.category_counts()
            top = ", ".join("%s=%d" % (c, n) for c, n in cats.most_common(5))
            needs_attention.append((name, top or "issues=%d" % issues))

    print("\n%d masters checked." % len(rows))
    print("\nClean candidates for just-in-time extraction (0 issues found):")
    if clean:
        for name in clean:
            print("  - %s" % name)
    else:
        print("  (none)")
    print("\nNeed human attention before/while extracting:")
    if needs_attention:
        for name, why in needs_attention:
            print("  - %-55s %s" % (name[:55], why))
    else:
        print("  (none)")
    return 0


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def main(argv):
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("master", nargs="?", help="path to one .docx master (mode 1)")
    ap.add_argument("--slug", help="asset/spec slug (default: derived from the filename)")
    ap.add_argument("--out", help="spec .json path (default: specs/<slug>.json)")
    ap.add_argument("--assets-out", dest="assets_out",
                    help="directory for extracted screenshots (default: "
                         "<brand assets_dir>/<slug>/)")
    ap.add_argument("--force", action="store_true",
                    help="overwrite an existing spec or asset file (default: refuse)")
    resolve_brand.add_brand_argument(ap)
    ap.add_argument("--report", action="store_true",
                    help="mode 2: report extractability of every master in the brand's masters_dir")
    args = ap.parse_args(argv[1:])

    global BRAND
    BRAND = resolve_brand.resolve(args.brand)
    S.configure(BRAND)

    # Console output in this environment can be a non-UTF-8 codepage; SOP text
    # routinely contains characters (em dashes, curly quotes) that codepage
    # can't encode. Never let a report crash on print() because of that.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass  # stdout is not reconfigurable in this environment; best effort

    if args.report:
        if args.master:
            ap.error("--report does not take a master path")
        return report_all()

    if not args.master:
        ap.error("a master .docx path is required unless --report is given")

    try:
        extract_one(args.master, slug=args.slug, out_path=args.out,
                   assets_dir=args.assets_out, force=args.force)
    except FileExistsError as exc:
        print("ERROR: %s" % exc, file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
