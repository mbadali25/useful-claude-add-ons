"""Built-in themes and densities: resolution order, identity protection,
contrast, dark-page carriage into the converted document, and the gallery.

SABOTAGE CHECKS (each was run; each turned this suite red):
  * In resolve_brand.Brand.apply_overlay, delete the `banned` pop loop ->
    test_a_theme_cannot_replace_brand_identity fails (a theme carrying a footer
    replaces Solomon's).
  * In any assets/themes/*.json, set table_head_ink equal to table_head ->
    test_every_theme_meets_contrast[...] fails for that theme.
  * Change one hex in a theme without rebuilding the gallery ->
    test_gallery_is_current fails.
"""

import json
import os
import shutil
import sys
import zipfile

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)

import build_gallery  # noqa: E402
import build_report  # noqa: E402
import house_style  # noqa: E402
import resolve_brand  # noqa: E402

SKILLS = os.path.normpath(os.path.join(SCRIPTS, os.pardir, os.pardir))
THEMES = [n for n, _, _ in resolve_brand.list_themes()]
EXPECTED_THEMES = {"professional", "corporate", "blue", "red", "modern",
                   "dark", "midnight", "high-contrast"}


@pytest.fixture(autouse=True)
def _clean_env(monkeypatch):
    for var in (resolve_brand.THEME_ENV, resolve_brand.DENSITY_ENV, resolve_brand.BRAND_ENV):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv(resolve_brand.SKILLS_DIR_ENV, SKILLS)


def _resolve(brand="neutral", theme=None, density=None):
    return resolve_brand.resolve(brand, announce=False, theme=theme, density=density)


def _lum(hexv):
    def ch(c):
        c = c / 255
        return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
    h = hexv.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast(a, b):
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def test_the_expected_themes_ship():
    """A superset check: adding a theme file must not fail the suite (SKILL.md
    says to add one by copying a JSON file); removing a shipped one must."""
    assert EXPECTED_THEMES <= set(THEMES)


def test_densities_ship():
    assert {n for n, _ in resolve_brand.list_densities()} == {"comfortable", "compact"}


@pytest.mark.parametrize("wanted", ["High Contrast", "high_contrast", "HIGH-CONTRAST"])
def test_theme_names_are_forgiving(wanted):
    assert _resolve(theme=wanted).theme == "high-contrast"


def test_unknown_theme_names_the_available_ones():
    with pytest.raises(resolve_brand.OverlayNotFound, match="midnight"):
        _resolve(theme="sparkly")


def _parse(argv):
    import argparse  # pylint: disable=import-outside-toplevel
    ap = argparse.ArgumentParser()
    resolve_brand.add_theme_arguments(ap)
    return ap.parse_args(argv)


def test_theme_from_environment_reaches_build_scripts(monkeypatch):
    monkeypatch.setenv(resolve_brand.THEME_ENV, "red")
    args = _parse([])
    assert _resolve(theme=args.theme).theme == "red"


def test_explicit_theme_beats_environment(monkeypatch):
    monkeypatch.setenv(resolve_brand.THEME_ENV, "red")
    assert _parse(["--theme", "blue"]).theme == "blue"


@pytest.mark.parametrize("var", ["THEME_ENV", "DENSITY_ENV"])
def test_resolve_itself_ignores_the_environment(monkeypatch, var):
    """extract_spec, the checkers and the gallery call resolve() without a
    theme; an overlay env var must not reach them."""
    monkeypatch.setenv(getattr(resolve_brand, var), "compact" if var == "DENSITY_ENV" else "red")
    b = _resolve()
    assert b.theme is None and b.density is None


def test_gallery_check_ignores_density_env(monkeypatch):
    monkeypatch.setenv(resolve_brand.DENSITY_ENV, "compact")
    assert build_gallery.main(["--check"]) == 0


def test_a_broken_theme_file_does_not_stop_builds_that_ignore_themes(tmp_path, monkeypatch, capsys):
    bad = tmp_path / "themes"
    bad.mkdir()
    for f in os.listdir(resolve_brand.THEMES_DIR):
        if f.endswith(".json"):
            shutil.copyfile(os.path.join(resolve_brand.THEMES_DIR, f), bad / f)
    (bad / "mine.json").write_text('{"report": {')
    monkeypatch.setattr(resolve_brand, "THEMES_DIR", str(bad))
    names = [n for n, _, _ in resolve_brand.list_themes()]
    assert "mine" not in names and "midnight" in names
    assert "mine.json" in capsys.readouterr().err
    with pytest.raises(resolve_brand.OverlayInvalid, match="mine.json"):
        _resolve(theme="mine")


def test_no_theme_keeps_the_brand_colours():
    solomon = _resolve("solomon")
    assert solomon.theme is None
    assert solomon.report["navy"] == "#0E2841"


def test_a_theme_wins_on_colour_over_the_brand():
    themed = _resolve("solomon", theme="red")
    red = resolve_brand.load_overlay("theme", "red")[1]
    assert themed.report["navy"] == red["report"]["navy"]


def test_a_theme_keeps_the_brand_identity():
    plain, themed = _resolve("solomon"), _resolve("solomon", theme="midnight")
    assert themed.fonts == plain.fonts
    assert themed.logo == plain.logo
    assert themed.sop["footer"] == plain.sop["footer"]
    assert themed.template == plain.template
    assert themed.organisation == plain.organisation


def test_a_theme_cannot_replace_brand_identity():
    brand = _resolve("solomon")
    footer = dict(brand.sop["footer"])
    brand.apply_overlay("theme", "rogue", {
        "sop": {"footer": {"left": "HIJACKED"}, "template": "x.docx", "accent": "123456"},
        "report": {"output_dir": "/elsewhere", "navy": "#123456"},
        "fonts": {"body": "Comic Sans MS"},
    })
    assert brand.sop["footer"] == footer
    assert brand.sop["accent"] == "123456"
    assert brand.report["output_dir"] != "/elsewhere"
    assert brand.report["navy"] == "#123456"
    assert brand.fonts["body"] != "Comic Sans MS"


def test_describe_names_theme_and_density():
    text = _resolve(theme="dark", density="compact").describe()
    assert "theme dark" in text and "density compact" in text


# --------------------------------------------------------------------------
# Contrast: every text/background pair a theme produces
# --------------------------------------------------------------------------

@pytest.mark.parametrize("theme", THEMES)
def test_every_theme_meets_contrast(theme):
    pal = house_style.Palette(_resolve(theme=theme))
    floor = 7.0 if theme == "high-contrast" else 4.5
    pairs = {
        "body text on page": (pal.ink, pal.page),
        "body text on row": (pal.ink, pal.row),
        "body text on zebra": (pal.ink, pal.zebra),
        "body text on panel": (pal.ink, pal.panel),
        "header text on table head": (pal.table_head_ink, pal.table_head),
        "title on band": (pal.title_ink, pal.navy),
        "org text on band": (pal.org_ink, pal.navy),
        "heading on page": (pal.heading, pal.page),
        "muted on page": (pal.muted, pal.page),
        "classification bar": ("#FFFFFF", pal.classification),
        "handling notice": (pal.handling_ink, pal.handling_bg),
        "card label (muted) on panel": (pal.muted, pal.panel),
        "card number (heading) on panel": (pal.heading, pal.panel),
        "guide .warn on page": (pal.warn_ink, pal.page),
        "guide .fail on page": (pal.fail_ink, pal.page),
        "link on page": (pal.link, pal.page),
    }
    bad = {k: round(contrast(*v), 2) for k, v in pairs.items() if contrast(*v) < floor}
    assert not bad, f"{theme} below {floor}:1 -> {bad}"


@pytest.mark.parametrize("theme", THEMES)
def test_every_theme_keeps_the_table_grid_visible(theme):
    """A border the same colour as the cell is no border. 1.3:1 is a floor
    for a visible non-text edge, well under WCAG's 3:1 for UI components on
    purpose: a black grid on a black page is the failure, not a subtle one."""
    pal = house_style.Palette(_resolve(theme=theme))
    assert contrast(pal.table_border, pal.row) >= 1.3


@pytest.mark.parametrize("theme", THEMES)
def test_sop_colours_meet_contrast(theme):
    b = _resolve(theme=theme)
    page = "#" + (b.sop.get("page") or "FFFFFF").lstrip("#")
    for key in ("body", "heading", "title", "link", "caption"):
        assert contrast("#" + b.sop[key].lstrip("#"), page) >= 4.5, f"{theme} sop.{key}"


# --------------------------------------------------------------------------
# Density
# --------------------------------------------------------------------------

def test_compact_tightens_the_table():
    css = house_style.stylesheet(house_style.Palette(_resolve(density="compact")), "report")
    assert "padding:3px 6px" in css and "font-size:11.5px" in css


def test_comfortable_equals_the_default():
    a = house_style.stylesheet(house_style.Palette(_resolve()), "report")
    b = house_style.stylesheet(house_style.Palette(_resolve(density="comfortable")), "report")
    assert a == b


def test_density_combines_with_any_theme():
    b = _resolve(theme="red", density="compact")
    red = resolve_brand.load_overlay("theme", "red")[1]
    assert b.report["navy"] == red["report"]["navy"]
    assert b.report["density"]["cell_pad"] == "3px 6px"


# --------------------------------------------------------------------------
# Dark pages survive conversion
# --------------------------------------------------------------------------

def test_white_page_report_is_unmarked():
    html = build_report.build(build_report.EXAMPLE, _resolve(theme="professional"))
    assert house_style.PAGE_META not in html and "bgcolor" not in html


@pytest.mark.parametrize("theme", ["midnight", "high-contrast"])
def test_dark_page_report_is_marked_three_ways(theme, tmp_path):
    pal = house_style.Palette(_resolve(theme=theme))
    html = build_report.build(build_report.EXAMPLE, _resolve(theme=theme))
    assert f"background:{pal.page}" in html
    assert f'<body bgcolor="{pal.page}">' in html
    assert f'<meta name="{house_style.PAGE_META}" content="{pal.page}">' in html
    src = tmp_path / "r.html"
    src.write_text(html, encoding="utf-8")
    h = pal.page.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    assert build_report.page_colour(str(src)) == r + (g << 8) + (b << 16)


def test_apply_to_html_marks_a_dark_guide():
    pal = house_style.Palette(_resolve(theme="midnight"))
    out = house_style.apply_to_html("<html><head><style></style></head><body><h1>x</h1></body></html>",
                                    pal, "guide")
    assert f'bgcolor="{pal.page}"' in out


@pytest.mark.parametrize("theme,dark", [("professional", False), ("midnight", True)])
def test_sop_page_colour(theme, dark, tmp_path):
    pytest.importorskip("docx")
    import build_sop  # pylint: disable=import-outside-toplevel
    brand = _resolve(theme=theme)
    b = build_sop.SopBuilder("Title", "Sub", brand=brand)
    b.para("Hello")
    out = b.save(str(tmp_path / "s.docx"))
    with zipfile.ZipFile(out) as z:
        doc = z.read("word/document.xml").decode()
        settings = z.read("word/settings.xml").decode()
    assert ("<w:background" in doc) is dark
    assert ("displayBackgroundShape" in settings) is dark
    if dark:
        order = _settings_order(settings)
        for before in build_report.SETTINGS_BEFORE_BG:
            if before in order:
                assert order.index("displayBackgroundShape") > order.index(before), before
    if dark:
        assert build_report.page_colour(out) is not None


# --------------------------------------------------------------------------
# Gallery
# --------------------------------------------------------------------------

def test_gallery_covers_every_theme():
    stems = {s for s, *_ in build_gallery.samples()}
    assert set(THEMES) <= stems and "professional-compact" in stems


def test_gallery_is_current():
    assert build_gallery.main(["--check"]) == 0


def test_every_theme_file_is_self_describing():
    for name in THEMES:
        _, data = resolve_brand.load_overlay("theme", name)
        assert data.get("label") and data.get("description"), name
        assert set(data) <= {"_comment", "name", "label", "description", "sources",
                             "report", "sop"}, name
        json.dumps(data)


@pytest.mark.parametrize("brand", ["neutral", "solomon"])
@pytest.mark.parametrize("theme", ["midnight", "high-contrast"])
def test_sop_header_footer_text_is_readable_on_a_dark_page(brand, theme, tmp_path):
    """A brand TEMPLATE brings its own footer runs (Solomon: 7F7F7F, 3.1:1 on
    midnight); on a dark page every header/footer run must meet the floor."""
    pytest.importorskip("docx")
    import re as _re  # pylint: disable=import-outside-toplevel
    import build_sop  # pylint: disable=import-outside-toplevel
    b = _resolve(brand, theme=theme)
    sop = build_sop.SopBuilder("Title", "Sub", brand=b)
    sop.para("Body")
    out = sop.save(str(tmp_path / "s.docx"))
    page = "#" + b.sop["page"].lstrip("#")
    floor = 7.0 if theme == "high-contrast" else 4.5
    with zipfile.ZipFile(out) as z:
        parts = [n for n in z.namelist() if _re.match(r"word/(header|footer)\d*\.xml$", n)]
        colours = set()
        for n in parts:
            colours |= set(_re.findall(r'<w:color w:val="([0-9A-Fa-f]{6})"', z.read(n).decode()))
    assert parts
    bad = {c: round(contrast("#" + c, page), 2) for c in colours if contrast("#" + c, page) < floor}
    assert not bad, bad


# --------------------------------------------------------------------------
# The LibreOffice DOCX page-colour stamp, and CT_Settings ordering
# --------------------------------------------------------------------------

def _settings_order(xml):
    import re as _re  # pylint: disable=import-outside-toplevel
    return _re.findall(r"<w:([A-Za-z]+)\b", xml.split("<w:settings", 1)[1])[0:]


def test_insert_display_bg_follows_every_child_the_schema_puts_first():
    xml = ('<w:settings xmlns:w="x"><w:view w:val="web"/><w:zoom w:percent="100"/>'
           '<w:removePersonalInformation/><w:defaultTabStop w:val="720"/></w:settings>')
    out = build_report._insert_display_bg(xml)
    order = _settings_order(out)
    assert order.index("displayBackgroundShape") > order.index("removePersonalInformation")
    assert order.index("displayBackgroundShape") < order.index("defaultTabStop")
    assert build_report._insert_display_bg(out) == out  # idempotent


def test_stamp_page_colour_writes_background_and_keeps_the_package(tmp_path):
    pytest.importorskip("docx")
    import docx  # pylint: disable=import-outside-toplevel
    path = tmp_path / "plain.docx"
    d = docx.Document()
    d.add_paragraph("x")
    d.save(str(path))
    os.chmod(path, 0o644)
    with zipfile.ZipFile(path) as z:
        before = z.namelist()
    build_report.stamp_page_colour(str(path), "2E3440")
    with zipfile.ZipFile(path) as z:
        assert z.namelist() == before  # same parts, same order ([Content_Types].xml first)
        doc = z.read("word/document.xml").decode()
        settings = z.read("word/settings.xml").decode()
    assert '<w:background w:color="2E3440"/>' in doc
    order = _settings_order(settings)
    assert "displayBackgroundShape" in order
    if "zoom" in order:
        assert order.index("displayBackgroundShape") > order.index("zoom")
    assert oct(os.stat(path).st_mode & 0o777) == oct(0o644)
    assert build_report.page_colour(str(path)) == 0x2E + (0x34 << 8) + (0x40 << 16)


def test_libreoffice_convert_stamps_a_dark_page_docx(tmp_path, monkeypatch):
    """convert() must stamp the page colour after a LibreOffice DOCX export
    (which drops it). The engine is stubbed; the stamp is what is under test."""
    pytest.importorskip("docx")
    import docx  # pylint: disable=import-outside-toplevel
    html = build_report.build(build_report.EXAMPLE, _resolve(theme="midnight"))
    src = tmp_path / "r.html"
    src.write_text(html, encoding="utf-8")

    def fake_soffice(path, want_docx=False, want_pdf=False):
        d = docx.Document()
        d.add_paragraph("x")
        d.save(os.path.splitext(path)[0] + ".docx")
        return 0

    import render_engine  # pylint: disable=import-outside-toplevel
    monkeypatch.setattr(render_engine, "choose_engine", lambda r: (render_engine.LIBREOFFICE, "test"))
    monkeypatch.setattr(render_engine, "to_soffice", fake_soffice)
    assert build_report.convert(str(src), want_docx=True) == 0
    assert build_report.page_colour(str(tmp_path / "r.docx")) is not None


@pytest.mark.parametrize("theme", ["professional", "midnight"])
def test_restyle_in_place_replaces_the_page_marking(theme):
    """A page themed midnight and then restyled must carry the NEW theme's
    marking only - not the old dark bgcolor and meta."""
    dark = house_style.Palette(_resolve(theme="midnight"))
    new = house_style.Palette(_resolve(theme=theme))
    base = "<html><head><style></style></head><body><h1>x</h1></body></html>"
    once = house_style.apply_to_html(base, dark, "guide")
    twice = house_style.apply_to_html(once, new, "guide")
    assert twice.count("bgcolor") == (0 if new.page.upper() == "#FFFFFF" else 1)
    assert twice.count(house_style.PAGE_META) == (0 if new.page.upper() == "#FFFFFF" else 1)


def test_page_colour_finds_the_meta_after_a_long_head(tmp_path):
    pal = house_style.Palette(_resolve(theme="midnight"))
    base = "<html><head><script>" + "x" * 20000 + "</script><style></style></head><body></body></html>"
    src = tmp_path / "g.html"
    src.write_text(house_style.apply_to_html(base, pal, "guide"), encoding="utf-8")
    assert build_report.page_colour(str(src)) is not None
