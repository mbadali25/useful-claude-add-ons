"""
GATE 1 - check SOP .docx masters against the active brand pack's house template.

This is the machine-readable form of references/template-spec.md. Run it after
editing or generating a SOP, and before committing a batch:

    python check_conformance.py "C:\\path\\My SOP.docx" -v
    python check_conformance.py                      # every master in the brand's masters_dir
    python check_conformance.py --brand neutral out\\selftest.docx
    python check_conformance.py --no-spec-check      # skip the spec/master drift check

Exit code is non-zero if any file fails, so it can gate a build.

Checks are deliberately structural (margins, fonts, colours, the accent bar,
the footer, screenshot borders, wp:effectExtent, hyperlink formatting and
targets). Every expected value comes from the brand pack, so the SAME gate
validates a Solomon master against Solomon and a neutral document against
neutral. The footer check asserts the organisation text MATCHES THE ACTIVE
BRAND -- a Solomon footer on a document built for another brand fails, and so
does a missing one. It also compares a master against its JSON build spec in
the brand's specs directory, when one exists, to catch hand-edits that leave a
committed spec lying about the document it is supposed to reproduce. It does
not check wording (beyond that spec comparison) or screenshots.

Gate 1 reads the .docx and cannot see layout. Gate 2 is verify_borders.py.
Third-party requirement: python-docx.
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys

import docx
from docx.oxml.ns import qn

import build_sop as S
import resolve_brand

TEXT_WIDTH_IN = 7.5

# Spec markup: "[label](url)" -> "label", "**bold**" -> "bold". Used both to
# strip a spec's text down to the visible text a reader (and python-docx)
# would see, and to count hyperlinks a spec declares.
_LINK_MD_RE = re.compile(r"\[([^\]]*)\]\([^)]*\)")


def _rgb(run):
    try:
        c = run.font.color
        return str(c.rgb) if c is not None and c.rgb is not None else None
    except (AttributeError, ValueError):
        return None


def _strip_markup(text):
    """Spec inline markup -> the visible text python-docx would read back."""
    return _LINK_MD_RE.sub(r"\1", text or "").replace("**", "")


def _classify_para(para):
    """Classify one body paragraph the way build_sop.py's block types
    render it, mirroring the heuristics already used above (heading font/size,
    the "N." step marker, the bullet character, caption size). Returns
    (block_type, comparable_text); comparable_text is None for blocks with no
    text of their own (images, spacers)."""
    if para._p.findall(".//" + qn("w:drawing")):
        return ("image", None)
    runs_with_text = [r for r in para.runs if r.text]
    if not runs_with_text:
        return ("spacer", None)
    first = para.runs[0]
    size_pt = first.font.size.pt if first.font.size else None
    text = para.text
    if first.font.name == S.HEADING_FONT and size_pt == S.HEADING_PT:
        return ("heading", text)
    marker = first.text.strip()
    if marker and re.match(r"^\d+\.$", marker):
        return ("step", re.sub(r"^\d+\.\t", "", text, count=1))
    if marker == S.BULLET_CHAR:
        return ("bullet", re.sub("^" + re.escape(S.BULLET_CHAR) + r"\t", "", text, count=1))
    if size_pt == S.CAPTION_PT:
        return ("caption", text)
    if all(r.bold for r in runs_with_text):
        return ("tip", text)
    return ("para", text)


def _spec_blocks(spec):
    """Expected (block_type, comparable_text) sequence for a spec's body,
    mirroring build_sop.build_from_spec: an "image" block with a caption
    renders as two master paragraphs (the image, then the caption)."""
    blocks = []
    for blk in spec["body"]:
        kind = blk["type"]
        if kind == "image":
            blocks.append(("image", None))
            cap = blk.get("caption", "")
            if cap:
                blocks.append(("caption", _strip_markup(cap)))
        elif kind == "spacer":
            blocks.append(("spacer", None))
        else:
            blocks.append((kind, _strip_markup(blk.get("text", ""))))
    return blocks


def _spec_link_count(spec):
    """Count of [label](url) hyperlink markup across a spec's text/caption
    fields, for comparison against the master's actual <w:hyperlink> count."""
    count = 0
    for blk in spec["body"]:
        text = blk.get("text") or blk.get("caption") or ""
        count += len(_LINK_MD_RE.findall(text))
    return count


def _diff_blocks(spec_name, expected, actual):
    """First concrete divergence between a spec's expected block sequence and
    a master's actual one, as a human-readable string -- or None if they
    match block-for-block (which also means their per-type counts match)."""
    for i in range(max(len(expected), len(actual))):
        if i >= len(expected):
            at, atext = actual[i]
            return (f"spec {spec_name} has only {len(expected):d} block(s), but master has an extra "
                     f"block {i:d}: {at} {atext!r}")
        if i >= len(actual):
            et, etext = expected[i]
            return (f"spec {spec_name} says block {i:d} is a {et} {etext!r}, but master has no "
                     f"block there (master has only {len(actual):d} block(s))")
        et, etext = expected[i]
        at, atext = actual[i]
        if et != at:
            return (f"spec {spec_name} says block {i:d} is a {et} "
                    f"{etext!r}, master has a {at} {atext!r}")
        if etext is not None and atext is not None and etext.strip() != atext.strip():
            return (f"spec {spec_name} says block {i:d} ({et}) text is "
                    f"{etext!r}, master has {atext!r}")
    return None


def _load_spec_map(specs_dir):
    """Map normalised master path -> spec .json path, built from every spec
    in the brand's specs directory whose filename does not start with "_"
    (those are self-test fixtures, not real SOPs). No directory, no map."""
    mapping = {}
    if not specs_dir or not os.path.isdir(specs_dir):
        return mapping
    for spec_path in glob.glob(os.path.join(specs_dir, "*.json")):
        if os.path.basename(spec_path).startswith("_"):
            continue
        try:
            with open(spec_path, encoding="utf-8") as fh:
                spec = json.load(fh)
        except (OSError, ValueError):
            continue
        out = spec.get("output")
        if not out:
            continue
        if not os.path.isabs(out):
            out = os.path.join(os.path.dirname(spec_path), out)
        key = os.path.normcase(os.path.normpath(os.path.abspath(out)))
        mapping[key] = spec_path
    return mapping


def check(path, spec_map=None, check_spec=True):
    """Return (errors, warnings, info) for one .docx master, against the brand
    build_sop is currently configured for (call S.configure(brand) first).

    spec_map, when given, is a {normalised master path: spec .json path} map
    from _load_spec_map() -- pass one in when checking a batch so it is built
    once. If check_spec is true and spec_map is None, one is built here.
    """
    if S.STYLE is None:
        S.configure()
    style = S.STYLE
    errors, warnings, info = [], [], {}
    d = docx.Document(path)
    s = d.sections[0]

    # -- page setup ---------------------------------------------------------
    house = style.MARGINS_IN
    margins = (round(s.top_margin.inches, 2), round(s.bottom_margin.inches, 2),
               round(s.left_margin.inches, 2), round(s.right_margin.inches, 2))
    m_top, m_bot, m_left, m_right = margins
    info["margins"] = (f"T{m_top:.2f} B{m_bot:.2f} "
                       f"L{m_left:.2f} R{m_right:.2f}")
    if margins[0] < house["top"]:
        # House consistency only. This was once recorded as the fix for clipped
        # screenshot borders -- it is not. See the effectExtent check below.
        warnings.append(f"top margin is {margins[0]:.2f}\", house is {house['top']:.2f}\"")
    if margins[1:] != (round(house["bottom"], 2), round(house["left"], 2), round(house["right"], 2)):
        warnings.append(
            f"bottom/left/right margins are "
            f"{m_bot:.2f}/{m_left:.2f}/{m_right:.2f}, house is "
            f'{house["bottom"]:.2f}/{house["left"]:.2f}/{house["right"]:.2f}')

    # -- base font ----------------------------------------------------------
    normal = d.styles["Normal"].font.name
    info["normal_font"] = normal
    if normal != S.BODY_FONT:
        errors.append(f"Normal style font is {normal!r}, house is {S.BODY_FONT!r}")

    # -- title banner -------------------------------------------------------
    if not d.tables:
        errors.append("no title banner table (the accent bar + title block is missing)")
    else:
        tbl = d.tables[0]
        xml = tbl._tbl.xml
        fills = re.findall(r'<w:shd[^>]*w:fill="([0-9A-Fa-f]{6})"', xml)
        if S.ACCENT_RED not in [f.upper() for f in fills]:
            errors.append(f"accent bar cell is not shaded {S.ACCENT_RED} (found {fills})")
        cells = tbl.rows[0].cells
        title_run = None
        for para in cells[-1].paragraphs:
            if para.runs:
                title_run = para.runs[0]
                break
        if title_run is None:
            errors.append("title banner has no title text")
        else:
            info["title"] = title_run.text
            if title_run.font.name != S.HEADING_FONT:
                errors.append(f"title font is {title_run.font.name!r}, house is {S.HEADING_FONT!r}")
            pt = title_run.font.size.pt if title_run.font.size else None
            if pt != S.TITLE_PT:
                errors.append(f"title size is {pt} pt, house is {S.TITLE_PT} pt")
            if _rgb(title_run) != S.TITLE_GREY:
                errors.append(f"title colour is {_rgb(title_run)}, house is {S.TITLE_GREY}")

    # -- footer -------------------------------------------------------------
    # Inverted from the original gate, which asserted one hard-coded
    # organisation. This asserts the footer carries the ACTIVE BRAND's text:
    # a Solomon footer on a neutral document is as wrong as no footer.
    required = list(style.FOOTER.get("required_text") or [])
    optional = list(style.FOOTER.get("optional_text") or [])
    if not s.footer.tables:
        errors.append(f"no footer table (expected: {' / '.join(required + optional)})")
    else:
        ftext = " ".join(c.text for c in s.footer.tables[0].rows[0].cells)
        for want in required:
            if want not in ftext:
                errors.append(f"footer is missing {want!r} required by the {style.brand.name} brand")
        for want in optional:
            if want not in ftext:
                warnings.append(f"footer is missing {want!r}")

    # -- body blocks --------------------------------------------------------
    headings = steps = bullets = captions = 0
    for para in d.paragraphs:
        if not para.runs:
            continue
        first = para.runs[0]
        colour = _rgb(first)
        size_pt = first.font.size.pt if first.font.size else None

        if first.font.name == S.HEADING_FONT and size_pt == S.HEADING_PT:
            headings += 1
            if colour != S.HEADING_NAVY:
                errors.append(f"heading {para.text[:40]!r} colour is {colour}, house is {S.HEADING_NAVY}")
        elif re.match(r"^\d+\.$", first.text.strip()) and first.text.strip():
            steps += 1
            if colour != S.ACCENT_RED:
                errors.append(f"step marker {first.text.strip()!r} colour is {colour}, house is {S.ACCENT_RED}")
        elif first.text.strip() == S.BULLET_CHAR:
            bullets += 1
            if colour != S.ACCENT_RED:
                errors.append(f"bullet marker colour is {colour}, house is {S.ACCENT_RED}")
        elif size_pt == S.CAPTION_PT:
            captions += 1
            if colour != S.CAPTION_GREY:
                errors.append(f"caption {para.text[:40]!r} colour is {colour}, house is {S.CAPTION_GREY}")
            # Captions are italic in the house style -- measured 2026-09-08 as
            # 47 italic caption paragraphs across 17 masters and zero plain
            # ones. The spec previously recorded caption size and colour but
            # not italic, so the generator produced plain captions and this
            # checker passed them. Checked here so that cannot recur.
            # Only a caption with visible text: several masters carry empty
            # spacer paragraphs whose single run happens to be caption-sized,
            # and those have no runs to be italic.
            texted = [r for r in para.runs if r.text.strip()]
            if texted and not any(r.italic for r in texted):
                errors.append(f"caption {para.text[:40]!r} is not italic; house captions are")

    info["blocks"] = f"{headings:d} headings, {steps:d} steps, {bullets:d} bullets, {captions:d} captions"

    # -- screenshots --------------------------------------------------------
    body_xml = d.element.body.xml
    pics = re.findall(r"<pic:spPr>.*?</pic:spPr>", body_xml, re.S)
    extents = [(int(a), int(b)) for a, b in
               re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"', body_xml)]
    info["images"] = len(extents)
    unbordered = sum(1 for p in pics if S.IMAGE_BORDER_RED not in p.upper())
    if unbordered:
        errors.append(f"{unbordered:d} of {len(pics):d} screenshots have no {S.IMAGE_BORDER_RED} border")

    # Word strokes the picture outline OUTSIDE wp:extent. Without wp:effectExtent
    # reserving that band, the stroke falls outside the line box and Word clips
    # it -- the top border vanishes whenever the image sits at the top of a page.
    # This is the real cause of "clipped borders"; the top margin is irrelevant.
    starved = 0
    for inline in d.element.body.findall(".//" + qn("wp:inline")):
        if "<a:ln" not in inline.xml:
            continue
        ee = inline.find(qn("wp:effectExtent"))
        if ee is None or any(int(ee.get(s) or 0) < S.EFFECT_EXTENT_MIN for s in "ltrb"):
            starved += 1
    if starved:
        errors.append(f"{starved:d} bordered screenshot(s) have no adequate wp:effectExtent "
                      "-- their borders will be clipped at page top (run "
                      "fix_effect_extent.py)")
    for cx, _cy in extents:
        if cx > S.Inches(TEXT_WIDTH_IN):
            errors.append(f"a screenshot is {cx / 914400:.2f}\" wide, wider than the {TEXT_WIDTH_IN:.1f}\" text column")

    # -- hyperlinks -----------------------------------------------------------
    # A survey found 24 hyperlinks across 16 of the 19 masters, and nothing
    # checked them: a dropped or mis-formatted link is invisible to both
    # gates (Gate 2 rasterises the PDF and can't see r:id relationships).
    hyperlinks = d.element.body.findall(".//" + qn("w:hyperlink"))
    info["links"] = len(hyperlinks)
    rels = d.part.rels
    for h in hyperlinks:
        label = "".join(t.text or "" for t in h.findall(".//" + qn("w:t")))[:40]
        rid = h.get(qn("r:id"))
        if rid is None:
            errors.append(f"hyperlink {label!r} has no r:id -- renders as dead text")
        else:
            rel = rels.get(rid)
            if rel is None:
                errors.append(f"hyperlink {label!r} r:id {rid} has no matching relationship "
                              "-- dead link")
            elif not rel.is_external or not rel.target_ref.lower().startswith(
                    ("http://", "https://", "mailto:")):
                errors.append(f"hyperlink {label!r} (r:id {rid}) does not resolve to an external "
                              f"URL (target={rel.target_ref!r}, external={rel.is_external}) -- dead link")
        for run in h.findall(qn("w:r")):
            rPr = run.find(qn("w:rPr"))
            color = underline = None
            if rPr is not None:
                c = rPr.find(qn("w:color"))
                color = c.get(qn("w:val")) if c is not None else None
                u = rPr.find(qn("w:u"))
                underline = u.get(qn("w:val")) if u is not None else None
            if color is None or color.upper() != S.LINK_BLUE:
                errors.append(f"hyperlink {label!r} run colour is {color}, house is {S.LINK_BLUE}")
            if underline != "single":
                errors.append(f"hyperlink {label!r} underline is {underline!r}, house is 'single'")

    # -- spec / master drift ---------------------------------------------------
    # Specs are meant to make a master reproducible, but nothing stops someone
    # hand-editing the .docx afterwards -- leaving a committed spec that lies
    # about the document. Compare structure (not bytes) when a spec exists;
    # most masters have none today, and that is not an error.
    if check_spec:
        if spec_map is None:
            spec_map = _load_spec_map(style.brand.specs_dir)
        key = os.path.normcase(os.path.normpath(os.path.abspath(path)))
        spec_path = spec_map.get(key)
        if spec_path is None:
            info["spec"] = "none"
        else:
            spec_name = os.path.basename(spec_path)
            info["spec"] = spec_name
            try:
                with open(spec_path, encoding="utf-8") as fh:
                    spec = json.load(fh)
            except (OSError, ValueError) as exc:
                errors.append(f"spec {spec_name} could not be read: {exc}")
            else:
                expected = _spec_blocks(spec)
                actual = [_classify_para(p) for p in d.paragraphs]
                diff = _diff_blocks(spec_name, expected, actual)
                if diff:
                    errors.append(diff)
                else:
                    expected_links = _spec_link_count(spec)
                    if expected_links != info["links"]:
                        errors.append(
                            f"spec {spec_name} expects {expected_links:d} hyperlink(s), master has {info['links']:d}")

    return errors, warnings, info


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="*", help="specific .docx files (default: all masters)")
    ap.add_argument("--verbose", "-v", action="store_true", help="show details for passing files")
    ap.add_argument("--no-spec-check", action="store_true",
                    help="skip the spec/master drift check (Check B), e.g. while a spec "
                         "is deliberately mid-edit")
    ap.add_argument("--specs-dir", help="override the brand pack's specs directory")
    resolve_brand.add_brand_argument(ap)
    args = ap.parse_args()

    brand = resolve_brand.resolve(args.brand)
    S.configure(brand)

    paths = args.paths or sorted(
        glob.glob(os.path.join(brand.require_masters_dir(), "*.docx"))
    )
    paths = [p for p in paths if not os.path.basename(p).startswith("~$")]
    if not paths:
        print("No .docx masters found.", file=sys.stderr)
        return 2

    check_spec = not args.no_spec_check
    spec_map = _load_spec_map(args.specs_dir or brand.specs_dir) if check_spec else None

    failed = 0
    for path in paths:
        name = os.path.basename(path)
        try:
            errors, warnings, info = check(path, spec_map=spec_map, check_spec=check_spec)
        except Exception as exc:  # a corrupt master should not stop the batch
            print(f"ERROR {name[:52]!s:<52} could not be read: {exc}")
            failed += 1
            continue

        status = "FAIL " if errors else ("WARN " if warnings else "PASS ")
        print(f"{status}{name}")
        if errors or warnings or args.verbose:
            for k, v in info.items():
                print(f"        {k!s:<12} {v}")
        for e in errors:
            print(f"      - {e}")
        for w in warnings:
            print(f"      ~ {w}")
        if errors:
            failed += 1

    print(f"\n{len(paths):d} checked, {failed:d} failed.")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
